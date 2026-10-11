"""Extract IRENA's country and technology-specific real after-tax WACC assumptions for 2025
(Renewable power generation costs in 2025, data file, sheet 'Table A2') to a tidy CSV.

Input:  data/raw/manual/IRENA_TEC_RPGC_in_2025_data_file_2026.xlsx (manual download, see MANIFEST.csv)
Output: data/processed/irena_wacc_rpgc2025_tableA2.csv (fractions, e.g. 0.0577 = 5.77%)
Run:    python scripts/04d_extract_irena_wacc.py
"""
import os

import openpyxl
import pandas as pd

H = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
SRC = os.path.join(H, "data", "raw", "manual", "IRENA_TEC_RPGC_in_2025_data_file_2026.xlsx")
OUT = os.path.join(H, "data", "processed", "irena_wacc_rpgc2025_tableA2.csv")
COLS = {"Onshore Wind": "onshore_wind", "Solar PV": "solar_pv", "Offshore Wind": "offshore_wind",
        "Concentrating Solar Power": "csp", "Hydropower": "hydropower", "Geothermal": "geothermal",
        "Bioenergy": "bioenergy"}

ws = openpyxl.load_workbook(SRC, data_only=True, read_only=False)["Table A2"]
title = " ".join(str(ws.cell(r, c).value) for r in range(1, 4) for c in range(1, 4) if ws.cell(r, c).value)
assert "real after-tax WACC assumptions for 2025" in title, title
hdr_row = next(r for r in range(1, 10) if ws.cell(r, 2).value == "Country")
hdr = {ws.cell(hdr_row, c).value: c for c in range(2, ws.max_column + 1) if ws.cell(hdr_row, c).value}
rows = []
for r in range(hdr_row + 1, ws.max_row + 1):
    country = ws.cell(r, hdr["Country"]).value
    if not country:
        continue
    row = {"country": str(country).strip()}
    for k, v in COLS.items():
        row[v] = float(ws.cell(r, hdr[k]).value)
    row["provenance"] = "IRENA (2026) RPGC in 2025 data file, sheet Table A2"
    rows.append(row)
df = pd.DataFrame(rows)
df.to_csv(OUT, index=False)
print(f"{len(df)} countries -> {OUT}")
print(df[df.country.isin(["Greece", "Israel", "Egypt", "Cyprus", "Türkiye"])][["country", "solar_pv", "onshore_wind"]]
      .to_string(index=False))
